"""Optional independent workbook validation with openpyxl; no Revit needed."""
from pathlib import Path
from math import isclose
import openpyxl

root = Path(__file__).resolve().parents[1] / 'obj' / 'SmokeTests' / 'Output'
book = openpyxl.load_workbook(root / 'Smoke_BOM.xlsx')
assert book.sheetnames[:3] == ['BOM', 'Elements', 'Notes']
assert len(book.sheetnames) == 4
assert book['BOM']['B2'].value == 'Mô hình thử nghiệm'
assert book['BOM'].max_row == 9
assert book['Elements'].max_row == 10
amounts = [book['BOM'].cell(r, 8).value for r in range(6, 10)]
assert all(isinstance(v, (int, float)) for v in amounts)
assert any(isclose(v, 3.58023) for v in amounts)
assert isclose(sum(amounts), 15.58023)
assert sum(book['BOM'].cell(r, 9).value for r in range(6, 10)) == 5
assert book['BOM'].freeze_panes == 'A6'
book.close()

book = openpyxl.load_workbook(root / 'Template_Output.xlsx')
assert book['BOM']['K6'].value == '=SUM(H6:H100)'
assert book['BOM']['H9'].style_id == 2
assert book['BOM'].row_dimensions[9].height == 25
assert {'BOM_Elements', 'BOM_Notes'} <= set(book.sheetnames)
book.close()

book = openpyxl.load_workbook(root / 'Text_Limits.xlsx')
assert book['BOM']['E6'].data_type == 's'
assert book['BOM']['E6'].value.startswith('=HYPERLINK')
assert len(book['BOM']['J6'].value) == 32767
assert book['BOM']['D6'].value == 'Vietnamese: Ống'
book.close()
book = openpyxl.load_workbook(root / 'Categories_Custom_Layout.xlsx')
assert len(book.sheetnames) == 6
assert all(len(name) <= 31 for name in book.sheetnames)
assert book['BOM']['A1'].value == 'Custom BOM title'
assert book['BOM']['A5'].value == 'Measured length'
assert all(book['BOM'].cell(5, c).value != 'System' for c in range(1, 11))
assert any(cell.value == '<varies>' for row in book['BOM'].iter_rows() for cell in row)
assert all(book[name].max_row == 6 for name in book.sheetnames[3:])
assert sum(book[name]['A6'].value for name in book.sheetnames[3:]) == 11
book.close()
print('PASS: independent openpyxl load, totals, category sheets, customized columns, styles, formulas and text limits.')
